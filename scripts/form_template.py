#!/usr/bin/env python3
"""表单模板逆向：从"用过的表单"（照片/扫描件）还原出**空白模板**，并把后填的字迹切出来。

为什么这样做（实测依据，2026-09-30 在一张真实表单照片上量的）：
    颜色(B−R) 印刷/手写都 ≈ −12；笔画粗细 6.0 vs 4.4–8.4 px 重叠；墨色深度、边缘锐度同样重叠
    ⇒ **像素级特征分不开**。唯一可靠信号是"**结构**"：落在模板印刷块/表格线上的墨是印的，
    其余是后填的。所以流程是"先学模板，再切字迹"，而不是硬编码 ROI。

通用性怎么保证（不再是"一次性脚本"）：
    learn  → 把任何一张表单的**结构**存成 profile JSON（页面几何 + 网格线 + 印刷文字块）
    render → 从 profile **重新排版**出干净的空白模板（docx / PDF），不依赖照片质感
    split  → 拿 profile 切任意一张同款表单：字迹层 + 空白层
    profile 是数据、代码是通用的；换表单只需重学一次，不用改代码。

用法（也接进了 office_cli.py 的 form-* 子命令）：
    python3 form_template.py learn  scan.jpg -o form.json
    python3 form_template.py render form.json -o blank.docx
    python3 form_template.py blank  scan.jpg form.json -o clean_blank.png
    python3 form_template.py extract scan.jpg form.json -o handwriting.png
"""
from __future__ import annotations

import csv
import io
import json
import shutil
import subprocess
import sys
from pathlib import Path

import cv2
import numpy as np

A4_W, A4_H = 1654, 2339          # 200 dpi 的 A4 像素尺寸
TESSERACT = shutil.which("tesseract")


# ── 1. 页面校正 ────────────────────────────────────────────────────────────

def rectify(img: np.ndarray, target: tuple[int, int] = (A4_W, A4_H)) -> tuple[np.ndarray, dict]:
    """找到纸张四角并透视校正成 A4 比例；找不到就原样返回（info 里说明）。"""
    g = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    g = cv2.GaussianBlur(g, (5, 5), 0)
    _, th = cv2.threshold(g, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    cnts, _ = cv2.findContours(th, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if not cnts:
        return img, {"rectified": False, "reason": "no contour"}
    c = max(cnts, key=cv2.contourArea)
    peri = cv2.arcLength(c, True)
    for eps in (0.02, 0.03, 0.05, 0.08):
        approx = cv2.approxPolyDP(c, eps * peri, True)
        if len(approx) == 4 and cv2.contourArea(approx) > 0.3 * img.shape[0] * img.shape[1]:
            pts = _order_quad(approx.reshape(4, 2).astype("float32"))
            (tl, tr, br, bl) = pts
            src = np.array([tl, tr, br, bl], "float32")
            tw = max(np.linalg.norm(br - bl), np.linalg.norm(tr - tl))
            th_ = max(np.linalg.norm(tr - br), np.linalg.norm(tl - bl))
            W, H = target if th_ / max(tw, 1) > 1.0 else (target[1], target[0])
            dst = np.array([[0, 0], [W - 1, 0], [W - 1, H - 1], [0, H - 1]], "float32")
            M = cv2.getPerspectiveTransform(src, dst)
            return cv2.warpPerspective(img, M, (W, H)), {"rectified": True, "size": [W, H]}
    return img, {"rectified": False, "reason": "paper quad not found"}


def _order_quad(p: np.ndarray) -> np.ndarray:
    s = p.sum(1)
    d = np.diff(p, axis=1).ravel()
    return np.array([p[np.argmin(s)], p[np.argmin(d)], p[np.argmax(s)], p[np.argmax(d)]], "float32")


# ── 2. 网格线 ─────────────────────────────────────────────────────────────

def _binarize(img: np.ndarray) -> np.ndarray:
    g = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    g = cv2.createCLAHE(3.0, (8, 8)).apply(g)
    g = cv2.GaussianBlur(g, (3, 3), 0)
    return cv2.adaptiveThreshold(g, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 41, 10)


def grid_lines(binimg: np.ndarray, min_ratio: float = 0.12, max_thick: int = 14) -> dict:
    """横/竖印刷直线。判据：又长又细（长度 ≥ 页宽 min_ratio、厚度 ≤ max_thick）。

    min_ratio 取得小、厚度放宽，是因为这类手机拍的表单里表格线很淡（实测按 0.22/9 只剩 5 条）。
    """
    h, w = binimg.shape
    out = {"h": [], "v": []}
    for axis, klen in (("h", int(w * min_ratio)), ("v", int(h * min_ratio))):
        k = cv2.getStructuringElement(cv2.MORPH_RECT, (klen, 1) if axis == "h" else (1, klen))
        opened = cv2.morphologyEx(binimg, cv2.MORPH_OPEN, k)
        n, lab, stats, _ = cv2.connectedComponentsWithStats(opened, 8)
        found = []
        for i in range(1, n):
            x, y, bw, bh = (stats[i, cv2.CC_STAT_LEFT], stats[i, cv2.CC_STAT_TOP],
                            stats[i, cv2.CC_STAT_WIDTH], stats[i, cv2.CC_STAT_HEIGHT])
            if axis == "h" and bh <= max_thick and bw >= klen:
                found.append((y + bh // 2, x, x + bw))
            elif axis == "v" and bw <= max_thick and bh >= klen:
                found.append((x + bw // 2, y, y + bh))
        found.sort()
        merged: list[list] = []
        for pos, a, b in found:                      # 合并邻近的重复检测
            if merged and pos - merged[-1][0] <= 4:
                m = merged[-1]
                m[0] = (m[0] + pos) // 2
                m[1], m[2] = min(m[1], a), max(m[2], b)
            else:
                merged.append([pos, a, b])
        out[axis] = [{"pos": int(p), "from": int(a), "to": int(b)} for p, a, b in merged]
    return out


# ── 3. OCR：印刷文字块 ────────────────────────────────────────────────────

def ocr_blocks(img: np.ndarray, lang: str = "chi_sim", min_conf: float = 30.0) -> list[dict]:
    """tesseract TSV → 词框；按行合并成文字块（模板的"印刷内容"就靠它认）。"""
    if TESSERACT is None:
        return []
    ok, buf = cv2.imencode(".png", img)
    if not ok:
        return []
    with io.BytesIO() as fh:
        fh.write(buf.tobytes())
        fh.seek(0)
        proc = subprocess.run([TESSERACT, "stdin", "stdout", "-l", lang, "tsv"],
                              input=fh.read(), capture_output=True, timeout=300)
    words = []
    for r in csv.DictReader(io.StringIO(proc.stdout.decode("utf-8", "replace")),
                            delimiter="\t", quoting=csv.QUOTE_NONE):
        try:
            if r["text"].strip() and float(r["conf"]) >= min_conf and int(r["height"]) > 5:
                words.append({"x": int(r["left"]), "y": int(r["top"]),
                              "w": int(r["width"]), "h": int(r["height"]), "t": r["text"]})
        except (KeyError, ValueError):
            continue
    words.sort(key=lambda d: (d["y"], d["x"]))
    if not words:
        return []
    med_h = float(np.median([wd["h"] for wd in words]))
    lines: list[dict] = []
    for wd in words:                                  # 按 y 中心聚成行（容差随字高走）
        cy = wd["y"] + wd["h"] / 2
        best, best_d = None, 1e9
        for ln in lines:
            d = abs(cy - ln["cy"])
            if d < best_d and d <= max(0.7 * med_h, 0.6 * ln["h"]):
                best, best_d = ln, d
        if best is None:
            lines.append({"cy": cy, "h": wd["h"], "y0": wd["y"], "y1": wd["y"] + wd["h"],
                          "x0": wd["x"], "x1": wd["x"] + wd["w"], "words": [wd]})
        else:
            best["words"].append(wd)
            best["cy"] = (best["cy"] * (len(best["words"]) - 1) + cy) / len(best["words"])
            best["h"] = max(best["h"], wd["h"])
            best["x0"], best["x1"] = min(best["x0"], wd["x"]), max(best["x1"], wd["x"] + wd["w"])
            best["y0"], best["y1"] = min(best["y0"], wd["y"]), max(best["y1"], wd["y"] + wd["h"])
    lines.sort(key=lambda d: d["y0"])
    return [{"bbox": [ln["x0"], ln["y0"], ln["x1"], ln["y1"]],
             "text": "".join(w["t"] for w in sorted(ln["words"], key=lambda d: d["x"]))}
            for ln in lines]


# ── 4. profile 学习 / 读写 ────────────────────────────────────────────────

def learn(img: np.ndarray, *, already_rectified: bool = False) -> dict:
    fixed, info = (img, {"rectified": None}) if already_rectified else rectify(img)
    binimg = _binarize(fixed)
    grid = grid_lines(binimg)
    blocks = ocr_blocks(fixed)
    return {
        "version": 1,
        "page": {"size": [fixed.shape[1], fixed.shape[0]], "rectified": info.get("rectified"),
                 "a4_ratio": round(fixed.shape[1] / fixed.shape[0], 4)},
        "grid": grid,
        "blocks": blocks,                 # 印刷文字块（模板内容）；可人工增删改
        "protect_margin": 5,
    }


def load_profile(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def prepare(img: np.ndarray, prof: dict) -> np.ndarray:
    """把任意一张同款扫描件对齐到 profile 的坐标系（先校正、再缩放到 profile 的页面尺寸）。

    必须做这一步：profile 的网格线/文字块坐标是在**校正后的页面**上量的；
    直接拿原图去套会整体错位（第一版就踩了这个坑，字迹一个都没去掉）。
    """
    fixed, _ = rectify(img)
    W, H = prof.get("page", {}).get("size", [fixed.shape[1], fixed.shape[0]])
    if (fixed.shape[1], fixed.shape[0]) != (W, H):
        fixed = cv2.resize(fixed, (W, H), interpolation=cv2.INTER_AREA)
    return fixed


def save_profile(prof: dict, path: str | Path) -> Path:
    path = Path(path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(prof, ensure_ascii=False, indent=1), encoding="utf-8")
    return path


# ── 5. 切分：模板墨 vs 后填墨 ─────────────────────────────────────────────

def protected_mask(shape: tuple[int, int], prof: dict, *, line_thick: int = 7) -> np.ndarray:
    """模板保护区 = 网格线 ＋ 印刷文字块（外扩 margin）。落在里面的墨＝印刷。"""
    m = np.zeros(shape, np.uint8)
    for ln in prof.get("grid", {}).get("h", []):
        m[max(ln["pos"] - line_thick, 0):ln["pos"] + line_thick + 1, ln["from"]:ln["to"]] = 255
    for ln in prof.get("grid", {}).get("v", []):
        m[ln["from"]:ln["to"], max(ln["pos"] - line_thick, 0):ln["pos"] + line_thick + 1] = 255
    mg = int(prof.get("protect_margin", 5))
    for blk in prof.get("blocks", []):
        x0, y0, x1, y1 = blk["bbox"]
        m[max(y0 - mg, 0):y1 + mg, max(x0 - mg, 0):x1 + mg] = 255
    return m


def protect_from_template(template_img: np.ndarray, shape: tuple[int, int], *,
                          margin: int = 6, thr: int = 150) -> np.ndarray:
    """用"重排出来的干净模板"当遮罩：模板上发暗的地方＝印刷内容。

    比 OCR 文字块可靠得多（OCR 漏一个字，那个印刷字就会被当手写抹掉）。通用做法：
    规格 → 干净模板（人可校对）→ 渲染成同尺寸图 → 取其暗像素当保护区。
    """
    g = cv2.cvtColor(template_img, cv2.COLOR_BGR2GRAY)
    if (g.shape[1], g.shape[0]) != (shape[1], shape[0]):
        g = cv2.resize(g, (shape[1], shape[0]), interpolation=cv2.INTER_AREA)
    m = (g < thr).astype(np.uint8) * 255
    return cv2.dilate(m, np.ones((margin, margin), np.uint8), 1)


def split_ink(img: np.ndarray, prof: dict, template_img: np.ndarray | None = None) -> dict:
    """返回 ink / handwriting / printed / protected 四张掩膜。

    给了 template_img（重排模板的渲染图）就用它当保护区；否则退回"网格线＋OCR 文字块"。
    """
    binimg = _binarize(img)
    lines = protected_mask(binimg.shape, {"grid": prof.get("grid", {}), "blocks": []}, line_thick=5)
    if prof.get("fill_regions"):
        # 「填表区」模式：只在这些矩形里找字迹——比"枚举所有印刷内容"稳得多，
        # 也是表单模板的常规做法（登记一次填写区，之后每张同款都能用）。
        regions = np.zeros(binimg.shape, np.uint8)
        for x0, y0, x1, y1 in prof["fill_regions"]:
            regions[max(y0, 0):y1, max(x0, 0):x1] = 255
        inside = cv2.bitwise_and(binimg, regions)
        ink = cv2.bitwise_and(inside, cv2.bitwise_not(lines))
        prot = cv2.bitwise_or(lines, cv2.bitwise_not(cv2.bitwise_or(regions, ink)))
    elif template_img is not None:
        prot = protect_from_template(template_img, binimg.shape)
        prot = cv2.bitwise_or(prot, lines)
        ink = cv2.bitwise_and(binimg, cv2.bitwise_not(prot))
    else:
        prot = protected_mask(binimg.shape, prof)
        ink = cv2.bitwise_and(binimg, cv2.bitwise_not(prot))
    return {"ink": binimg, "handwriting": ink, "printed": cv2.bitwise_and(binimg, prot), "protected": prot}


def to_layer(mask: np.ndarray, size=None, pad: int = 20) -> np.ndarray:
    """把掩膜变成白底黑字的图（可选裁掉空白）。"""
    img = 255 - mask
    if size == "tight":
        ys, xs = np.where(mask > 0)
        if len(xs):
            y0, y1 = max(ys.min() - pad, 0), min(ys.max() + pad, mask.shape[0])
            x0, x1 = max(xs.min() - pad, 0), min(xs.max() + pad, mask.shape[1])
            img = img[y0:y1, x0:x1]
    return img


def grow_regions(img: np.ndarray, prof: dict, *, pad: int = 45, min_area: int = 100,
                 max_area: int = 4000, max_w: int = 150, max_h: int = 100,
                 rounds: int = 2) -> dict:
    """自动"长大"填写区：找漏掉的手写字（在现有填写区附近、体积像字、不像整行印刷文字）。

    为什么要它：填写区是量出来的，第一次总有偏差（实测漏了"40"和日期数字）。
    做法是闭环——出一次空白 → 看还没抹掉的墨 → 把它们并进填写区 → 再出一版。
    """
    prof = json.loads(json.dumps(prof, ensure_ascii=False))       # 深拷贝
    regions = prof.setdefault("fill_regions", [])
    for _ in range(rounds):
        parts = split_ink(img, prof)
        near = np.zeros(parts["ink"].shape, np.uint8)
        for x0, y0, x1, y1 in regions:
            near[max(y0 - pad, 0):y1 + pad, max(x0 - pad, 0):x1 + pad] = 255
        resid = cv2.bitwise_and(parts["printed"], near)           # 填写区附近"被当成印刷保下来"的墨
        n, lab, st, _ = cv2.connectedComponentsWithStats(resid, 8)
        added = 0
        for i in range(1, n):
            a = st[i, cv2.CC_STAT_AREA]
            w_, h_ = st[i, cv2.CC_STAT_WIDTH], st[i, cv2.CC_STAT_HEIGHT]
            if min_area <= a <= max_area and w_ <= max_w and h_ <= max_h:
                bx, by = st[i, cv2.CC_STAT_LEFT], st[i, cv2.CC_STAT_TOP]
                regions.append([int(bx) - 4, int(by) - 4, int(bx + w_) + 4, int(by + h_) + 4])
                added += 1
        if not added:
            break
    return prof


def blank_raster(img: np.ndarray, prof: dict, *, binarize: bool = True,
                 template_img: np.ndarray | None = None) -> np.ndarray:
    """去掉后填字迹的干净空白：用大核中值估纸张底色填掉字迹，再把模板印刷内容/网格线贴回，最后压平光照。

    template_img 给"重排模板的渲染图"时最准（印刷字一个不漏）；不给则用 profile 的网格线＋OCR 块。
    """
    parts = split_ink(img, prof, template_img)
    erase = cv2.dilate(parts["handwriting"], np.ones((9, 9), np.uint8), 1)
    paper = cv2.medianBlur(img, 101)
    out = np.where(erase[..., None] > 0, paper, img)
    lines = protected_mask(img.shape[:2], prof, line_thick=4)
    out[lines > 0] = img[lines > 0]
    f = out.astype(np.float32)                       # 光照压平
    for c in range(3):
        bg = cv2.GaussianBlur(f[..., c], (0, 0), 25)
        f[..., c] = np.clip(f[..., c] / np.maximum(bg, 1) * 248, 0, 255)
    out = f.astype(np.uint8)
    if binarize:                                     # 像干净的复印件：Otsu 对"大片白+少量字"最稳，不产椒盐
        g = cv2.cvtColor(out, cv2.COLOR_BGR2GRAY)
        _, bw = cv2.threshold(g, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        out = cv2.cvtColor(bw, cv2.COLOR_GRAY2BGR)
    return out


# ── 6b. 规格驱动的模板渲染（通用：任何表单都给同一套渲染器）─────────────
#
# 为什么不用"OCR 自动重建"当主路：实测同一张照片，tesseract 把"产权人签字（盖章）"读成
# "产权人签字章)"、把多列文字串行，用它重建出来的模板是错的。**模板是标准件、文字是固定的**，
# 所以正确做法是：规格（JSON、数据）由人确认一次，渲染器（代码、通用）负责排版。

def _cm_to_dxa(cm: float) -> int:
    return int(round(cm * 1440 / 2.54))


def _set_spacing(p, d: dict) -> None:
    """行距/段前段后。``line_pt`` 走**固定行距**，这是把表格顶边钉到指定 cm 的唯一可靠办法
    （用默认行距时 Word 会按字体度量自行伸缩，实测比目标高 0.55cm）。"""
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt

    pf = p.paragraph_format
    if d.get("line_pt"):
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(d["line_pt"])
    if d.get("space_before_pt") is not None:
        pf.space_before = Pt(d["space_before_pt"])
    if d.get("space_after_pt") is not None:
        pf.space_after = Pt(d["space_after_pt"])
    if d.get("space_after") is not None:
        pf.space_after = Pt(d["space_after"])


def _set_cell_margins(cell, pad_cm) -> None:
    """单元格内边距。默认左右各 0.19cm，窄列里这点内边距足以把「住所（经营场所）」挤断行。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    if not pad_cm:
        return
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcMar"))
    if old is not None:
        tcPr.remove(old)
    mar = OxmlElement("w:tcMar")
    for tag in ("w:top", "w:start", "w:bottom", "w:end"):
        el = OxmlElement(tag)
        el.set(qn("w:w"), str(_cm_to_dxa(pad_cm)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _fix_table_layout(table, widths_cm: list[float]) -> None:
    """把列宽**钉死**。

    坑（2026-09-30 实测）：只写 ``cell.width = Cm(...)``，LibreOffice/Word 仍按内容自动
    适配 —— 渲染出来分隔线跑到 11.05cm（要的是 5.30cm）。必须同时给
    ``tblLayout=fixed`` + ``tblW`` + **重建 tblGrid 的 gridCol**，三样缺一不可。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm

    table.autofit = False
    tblPr = table._tbl.tblPr
    for tag in ("w:tblLayout", "w:tblW"):
        old = tblPr.find(qn(tag))
        if old is not None:
            tblPr.remove(old)
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:w"), str(_cm_to_dxa(sum(widths_cm))))
    tw.set(qn("w:type"), "dxa")
    tblPr.append(tw)

    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for gc in list(grid):
            grid.remove(gc)
        for wcm in widths_cm:
            gc = OxmlElement("w:gridCol")
            gc.set(qn("w:w"), str(_cm_to_dxa(wcm)))
            grid.append(gc)

    for row in table.rows:
        cells = row.cells
        for ci, wcm in enumerate(widths_cm):
            if ci < len(cells):
                cells[ci].width = Cm(wcm)


def render_spec(spec: dict | str | Path, out: str | Path) -> Path:
    """按规格渲染空白模板 docx。

    spec = {
      "page": {"margins_cm": {"top":2.2,"bottom":2.0,"left":2.5,"right":2.0}},
      "title": {"text": "...", "font": "黑体", "size": 22, "bold": true},
      "body": [
        {"table": {"widths_cm": [3.4, 11.2],
                   "rows": [{"height_cm": 2.4,
                             "cells": [{"text": "住所（经营场所）\\n（详细地址）", "bold": true, "align": "center"},
                                       {"text": ""}]}]}},
        {"p": {"text": "须知", "bold": true, "size": 14, "align": "center"}},
        {"p": {"text": "1、...", "size": 10.5, "indent_cm": 0.7}}
      ]}
    """
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Mm, Pt

    if isinstance(spec, (str, Path)):
        spec = json.loads(Path(spec).read_text(encoding="utf-8"))
    ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER, "left": WD_ALIGN_PARAGRAPH.LEFT,
             "right": WD_ALIGN_PARAGRAPH.RIGHT}

    def style_run(run, *, font="Noto Serif CJK SC", size=10.5, bold=False):
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        rpr = run._element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(a), font)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    mg = (spec.get("page") or {}).get("margins_cm", {})
    sec.top_margin, sec.bottom_margin = Cm(mg.get("top", 2.2)), Cm(mg.get("bottom", 2.0))
    sec.left_margin, sec.right_margin = Cm(mg.get("left", 2.5)), Cm(mg.get("right", 2.0))

    t = spec.get("title")
    if t:
        p = doc.add_paragraph()
        p.alignment = ALIGN.get(t.get("align", "center"), WD_ALIGN_PARAGRAPH.CENTER)
        _set_spacing(p, t)
        style_run(p.add_run(t["text"]), font=t.get("font", "Noto Serif CJK SC"),
                  size=t.get("size", 22), bold=t.get("bold", True))

    for el in spec.get("body", []):
        if "p" in el:
            d = el["p"]
            p = doc.add_paragraph()
            p.alignment = ALIGN.get(d.get("align", "left"), WD_ALIGN_PARAGRAPH.LEFT)
            pf = p.paragraph_format
            pf.space_after = Pt(d.get("space_after", 2))
            if d.get("indent_cm"):
                pf.left_indent = Cm(d["indent_cm"])
            if d.get("first_line_cm"):
                pf.first_line_indent = Cm(d["first_line_cm"])
            _set_spacing(p, d)
            style_run(p.add_run(d.get("text", "")), font=d.get("font", "Noto Serif CJK SC"),
                      size=d.get("size", 10.5), bold=d.get("bold", False))
        elif "table" in el:
            d = el["table"]
            rows = d.get("rows", [])
            if not rows:
                continue
            ncol = max(len(r.get("cells", [])) for r in rows)
            table = doc.add_table(rows=len(rows), cols=ncol)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, r in enumerate(rows):
                if r.get("height_cm"):
                    table.rows[ri].height = Cm(r["height_cm"])
                if r.get("height_rule"):
                    table.rows[ri].height_rule = {"exact": WD_ROW_HEIGHT_RULE.EXACTLY,
                                                  "atLeast": WD_ROW_HEIGHT_RULE.AT_LEAST}[r["height_rule"]]
                for ci, c in enumerate(r.get("cells", [])):
                    cell = table.cell(ri, ci)
                    cell.text = ""
                    if c.get("valign"):
                        cell.vertical_alignment = {"top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
                                                   "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
                                                   "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}[c["valign"]]
                    _set_cell_margins(cell, c.get("padding_cm"))
                    paras = c.get("paras") or [{"text": t} for t in str(c.get("text", "")).split("\n")]
                    for li, ln in enumerate(paras):
                        p = cell.paragraphs[0] if li == 0 else cell.add_paragraph()
                        p.alignment = ALIGN.get(ln.get("align", c.get("align", "left")), WD_ALIGN_PARAGRAPH.LEFT)
                        p.paragraph_format.space_after = Pt(ln.get("space_after", 0))
                        if ln.get("first_line_cm"):
                            p.paragraph_format.first_line_indent = Cm(ln["first_line_cm"])
                        _set_spacing(p, ln)
                        style_run(p.add_run(ln.get("text", "")),
                                  font=ln.get("font", c.get("font", "Noto Serif CJK SC")),
                                  size=ln.get("size", c.get("size", 10.5)),
                                  bold=ln.get("bold", c.get("bold", False)))
                    if c.get("merge_right") and ci + 1 < ncol:      # 横向合并
                        table.cell(ri, ci).merge(table.cell(ri, min(ci + int(c["merge_right"]), ncol - 1)))
            if d.get("widths_cm"):
                _fix_table_layout(table, d["widths_cm"])
        elif "pagebreak" in el:
            doc.add_page_break()

    out = Path(out).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


# ── 7. 出图自检：渲染 → PDF → 量线位 → 比目标 ─────────────────────────────

def measure_lines(pdf: str | Path) -> dict:
    """量出 PDF 里表格的横/竖线位置（cm，从纸左上角算）。用来验证排版是否落在目标几何上。

    只认「细长」图形：高 < 2pt 且宽 > 30pt 的横线、宽 < 2pt 且高 > 30pt 的竖线。
    表格边框在 PDF 里就是这样的矢量矩形/线段，比截图再找线可靠得多。
    """
    import pymupdf

    doc = pymupdf.open(str(pdf))
    CM = 72 / 2.54
    h, v = [], []
    npages = doc.page_count
    for page in doc:
        for dr in page.get_drawings():
            r = dr["rect"]
            if r.height < 2 and r.width > 30:
                h.append(round(r.y0 / CM, 3))
            elif r.width < 2 and r.height > 30:
                v.append(round(r.x0 / CM, 3))
    doc.close()

    def merge(vals: list[float], tol: float = 0.06) -> list[float]:
        out: list[list[float]] = []
        for x in sorted(vals):
            if out and x - out[-1][-1] <= tol:
                out[-1].append(x)
            else:
                out.append([x])
        return [round(sum(g) / len(g), 2) for g in out]

    return {"pages": npages, "h_cm": merge(h), "v_cm": merge(v)}


def verify_spec(spec: dict | str | Path, out: str | Path | None = None,
                keep_pdf: str | Path | None = None) -> dict:
    """按 spec 渲染 → 转 PDF → 量线位 → 和 ``spec["verify"]`` 比。

    spec["verify"] = {"pages": 1, "h_cm": [...], "v_cm": [...], "tol_cm": 0.15}

    为什么要它：docx 排版是"声明式"的，列宽会被自动适配、行高会被字体度量顶开、
    内容会溢出到第二页 —— 这些**只有渲染出来量一遍**才知道。踩过的坑：
    列宽失效（分隔线 11.05cm vs 目标 5.30）、整体上移 0.55cm、须知行溢出到第 2 页。
    """
    spec = _load_json_if_path(spec)
    exp = spec.get("verify") or {}
    tol = float(exp.get("tol_cm", 0.15))

    out = Path(out or "verify.docx").expanduser()
    render_spec(spec, out)
    pdf = Path(keep_pdf).expanduser() if keep_pdf else out.with_suffix(".pdf")
    _convert_to_pdf(out, pdf)
    got = measure_lines(pdf)

    def cmp(name, exp_list, got_list):
        rows, worst = [], 0.0
        for i, e in enumerate(exp_list):
            g = got_list[i] if i < len(got_list) else None
            d = abs(g - e) if g is not None else 999.0
            worst = max(worst, d)
            rows.append({"i": i, "target": e, "got": g, "delta": round(g - e, 2) if g is not None else None,
                         "ok": d <= tol})
        return rows, worst

    h_rows, h_worst = cmp("h", exp.get("h_cm", []), got["h_cm"])
    v_rows, v_worst = cmp("v", exp.get("v_cm", []), got["v_cm"])
    pages_ok = got["pages"] == exp.get("pages", got["pages"])
    extra_h = len(got["h_cm"]) - len(exp.get("h_cm", []))
    extra_v = len(got["v_cm"]) - len(exp.get("v_cm", []))
    ok = pages_ok and h_worst <= tol and v_worst <= tol and extra_h <= 0 and extra_v <= 0
    return {"ok": ok, "pages": got["pages"], "pages_expected": exp.get("pages"),
            "tol_cm": tol, "h": h_rows, "v": v_rows,
            "h_max_delta": round(h_worst, 3), "v_max_delta": round(v_worst, 3),
            "extra_h_lines": extra_h, "extra_v_lines": extra_v,
            "docx": str(out), "pdf": str(pdf)}


def _load_json_if_path(spec):
    if isinstance(spec, (str, Path)):
        return json.loads(Path(spec).read_text(encoding="utf-8"))
    return spec


def _convert_to_pdf(src: Path, dst: Path) -> None:
    """调 LibreOffice 转 PDF。每个进程独立 UserInstallation，避免并发抢 profile。"""
    import tempfile
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("没找到 soffice，无法出 PDF 自检")
    dst.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        cmd = [soffice, "--headless", "--norestore", f"-env:UserInstallation=file://{td}",
               "--convert-to", "pdf", "--outdir", str(dst.parent), str(src)]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        made = dst.parent / (src.stem + ".pdf")
        if not made.exists():
            raise SystemExit(f"LO 没出 PDF：{r.stdout[-400:]}{r.stderr[-400:]}")
        if made != dst:
            made.replace(dst)


# ── 8. 重新排版空白模板（从扫描 profile，近似）──────────────────────────

def render_docx(prof: dict, out: str | Path) -> Path:
    """按 profile 重排一份**干净的**空白模板（网格线 → 表格，印刷文字块 → 单元格文字）。

    定位方式：把检测到的竖线当作列边界、横线当作行边界，把每个文字块放进它所在的格子。
    这是"近似复原"：字号/精确间距按格子归一化，交付后建议人工扫一眼。
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    W, H = prof["page"]["size"]
    vx = sorted({l["pos"] for l in prof["grid"]["v"]}) or [int(W * 0.05), int(W * 0.95)]
    hy = sorted({l["pos"] for l in prof["grid"]["h"]}) or [int(H * 0.05), int(H * 0.95)]
    if vx[0] > W * 0.02:
        vx = [0] + vx
    if hy[0] > H * 0.02:
        hy = [0] + hy
    if vx[-1] < W * 0.98:
        vx = vx + [W]
    if hy[-1] < H * 0.98:
        hy = hy + [H]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    scale = 210.0 / W                                  # px → mm
    sec.top_margin = Mm(hy[0] * scale)
    sec.bottom_margin = Mm((H - hy[-1]) * scale)
    sec.left_margin = Mm(vx[0] * scale)
    sec.right_margin = Mm((W - vx[-1]) * scale)

    for ri in range(len(hy) - 1):
        for ci in range(len(vx) - 1):
            x0, x1 = vx[ci], vx[ci + 1]
            y0, y1 = hy[ri], hy[ri + 1]
            if x1 - x0 < W * 0.03 or y1 - y0 < 12:
                continue
            cell_blocks = [b for b in prof["blocks"]
                           if x0 - 8 <= b["bbox"][0] and b["bbox"][2] <= x1 + 8
                           and y0 - 8 <= b["bbox"][1] and b["bbox"][3] <= y1 + 8]
            if not cell_blocks and (x1 - x0) < W * 0.15:
                continue                                # 空格子且窄：多半是标签列，跳过
            t = doc.add_table(rows=1, cols=1)
            t.style = "Table Grid"
            cell = t.cell(0, 0)
            cell.text = ""
            for i, b in enumerate(sorted(cell_blocks, key=lambda d: d["bbox"][1])):
                h_px = max(b["bbox"][3] - b["bbox"][1], 1)
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                run = p.add_run(b["text"])
                run.font.name = "Noto Serif CJK SC"
                run.font.size = Pt(max(6.5, min(18.0, h_px * 72 / 200)))   # 200dpi ⇒ px→pt
            doc.add_paragraph()
    out = Path(out).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out


# ── CLI ──────────────────────────────────────────────────────────────────

def _read_image(p: str | Path) -> np.ndarray:
    img = cv2.imread(str(p))
    if img is None:
        raise SystemExit(f"读不到图：{p}")
    return img


def main(argv: list[str] | None = None) -> int:
    import argparse
    ap = argparse.ArgumentParser(prog="form_template.py", description="表单模板逆向：学模板 / 出空白 / 切字迹")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("learn"); p.add_argument("image"); p.add_argument("-o", "--out", required=True)
    p = sub.add_parser("render"); p.add_argument("profile"); p.add_argument("-o", "--out", required=True)
    p = sub.add_parser("render-spec"); p.add_argument("spec"); p.add_argument("-o", "--out", required=True)
    p = sub.add_parser("verify", help="渲染→出PDF→量线位→和 spec.verify 比")
    p.add_argument("spec"); p.add_argument("-o", "--out", help="docx 输出路径")
    p.add_argument("--keep-pdf", help="PDF 留在指定路径（默认放 docx 旁边）")
    p.add_argument("--tol", type=float, help="覆盖 spec 里的 tol_cm")
    for name in ("blank", "extract"):
        q = sub.add_parser(name); q.add_argument("image"); q.add_argument("profile"); q.add_argument("-o", "--out", required=True)
        q.add_argument("--tight", action="store_true")
        q.add_argument("--template", help="重排模板渲染图（PNG）")
    a = ap.parse_args(argv)
    if a.cmd == "learn":
        img, info = rectify(_read_image(a.image))
        prof = learn(img, already_rectified=True)
        prof["page"]["rectify"] = info
        print(json.dumps({"ok": True, "out": str(save_profile(prof, a.out)),
                          "blocks": len(prof["blocks"]),
                          "grid_h": len(prof["grid"]["h"]), "grid_v": len(prof["grid"]["v"])},
                         ensure_ascii=False))
    elif a.cmd == "render":
        print(json.dumps({"ok": True, "out": str(render_docx(load_profile(a.profile), a.out))}, ensure_ascii=False))
    elif a.cmd == "render-spec":
        spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
        print(json.dumps({"ok": True, "out": str(render_spec(spec, a.out))}, ensure_ascii=False))
    elif a.cmd == "verify":
        spec = json.loads(Path(a.spec).read_text(encoding="utf-8"))
        if a.tol:
            spec.setdefault("verify", {})["tol_cm"] = a.tol
        out = a.out or str(Path(a.spec).with_suffix(".verify.docx"))
        r = verify_spec(spec, out, getattr(a, "keep_pdf", None))
        print(json.dumps(r, ensure_ascii=False))
    elif a.cmd == "blank":
        prof = load_profile(a.profile)
        img = prepare(cv2.imread(str(a.image)), prof)
        out = Path(a.out).expanduser(); out.parent.mkdir(parents=True, exist_ok=True)
        tpl = cv2.imread(a.template) if getattr(a, "template", None) else None
        cv2.imwrite(str(out), blank_raster(img, prof, template_img=tpl))
        print(json.dumps({"ok": True, "out": str(out)}, ensure_ascii=False))
    elif a.cmd == "extract":
        prof = load_profile(a.profile)
        img = prepare(cv2.imread(str(a.image)), prof)
        tpl = cv2.imread(a.template) if getattr(a, "template", None) else None
        layer = to_layer(split_ink(img, prof, tpl)["handwriting"], "tight" if a.tight else None)
        out = Path(a.out).expanduser(); out.parent.mkdir(parents=True, exist_ok=True)
        cv2.imwrite(str(out), layer)
        print(json.dumps({"ok": True, "out": str(out)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
