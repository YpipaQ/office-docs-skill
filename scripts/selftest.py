#!/usr/bin/env python3
"""端到端自检：本机真跑一遍，不 mock。

覆盖：环境自检 → 生成 docx/xlsx/pptx → LO 转 PDF/CSV → 抽文本 → 替换 →
缩略图 → PDF 合并/拆分/渲染 → OCR（图片型 PDF）→ PDF 转 Word。

判据：每一步都断言**内容**（不是只看文件存在），最后打印 PASS/FAIL 汇总；
有任何 FAIL 退出码 1。
"""
from __future__ import annotations

import json
import shutil
import sys
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import cn_profiles            # noqa: E402
import convert as lo          # noqa: E402
import extras                 # noqa: E402
import pdf_ops                # noqa: E402
import readers                # noqa: E402
import writers                # noqa: E402

RESULTS: list[tuple[str, bool, str]] = []
CN = "中文自检标题"


def check(name: str, fn):
    try:
        detail = fn()
        RESULTS.append((name, True, str(detail or "")))
        print(f"PASS  {name}  {detail or ''}")
    except Exception as exc:
        RESULTS.append((name, False, f"{type(exc).__name__}: {exc}"))
        print(f"FAIL  {name}  {type(exc).__name__}: {exc}")
        traceback.print_exc(limit=3)


def main() -> int:
    work = Path(tempfile.mkdtemp(prefix="office-docs-selftest-"))

    # 1) 环境
    def env_ok():
        d = lo.doctor()
        assert d["ok"], f"环境不全：缺 {d['missing_libs']}，soffice={d['soffice']!r}"
        return f"soffice={Path(d['soffice']).name} cjk_fonts={d['cjk_fonts']}"
    check("doctor", env_ok)

    # 2) 生成 docx -> pdf -> 抽文本
    docx = work / "a.docx"
    spec = {"title": CN, "blocks": [
        {"h": 1, "text": "第一章 标题"},
        {"p": "段落：Hello Office 12345"},
        {"bullets": ["要点一", "要点二"]},
        {"table": {"rows": [["列A", "列B"], ["甲", "1"], ["乙", "2"]]}},
    ]}
    check("docx-new", lambda: f"{writers.new_docx(spec, docx).stat().st_size} B")

    def docx_to_pdf_text():
        pdf = lo.convert(docx, "pdf", work)
        txt = readers.text_of(pdf)
        for want in (CN, "第一章", "甲", "12345"):
            assert want in txt, f"PDF 里找不到 {want!r}"
        return f"{pdf.name} 取回 {len(txt)} 字符（中文/表格/数字都在）"
    check("convert docx->pdf + 抽文本", docx_to_pdf_text)

    def docx_replace():
        r = writers.docx_replace(docx, {"12345": "67890", "甲": "AAA"}, work / "a-rep.docx")
        assert r["hits"] >= 2, r
        txt = readers.text_of(work / "a-rep.docx")
        assert "67890" in txt and "AAA" in txt and "12345" not in txt, txt[:200]
        return f"{r['hits']} 处命中"
    check("docx-replace", docx_replace)

    # 3) xlsx 公式 -> LO 重算 -> CSV
    def xlsx_recalc():
        x = writers.new_xlsx({"sheets": [{"name": "数据", "header": True, "widths": [14, 10],
                                          "rows": [["项目", "值"], ["单价", 3.5], ["数量", 4],
                                                   ["合计", "=B2*B3"]]}]}, work / "b.xlsx")
        csv_out = lo.convert(x, "csv", work)     # LO 导出时会算公式
        rows = readers.read_table(csv_out)
        flat = json.dumps(rows, ensure_ascii=False)
        assert "14" in flat, flat
        assert "合计" in flat, flat
        return f"公式 B2*B3 = 14 已算出（{csv_out.name}）"
    check("xlsx-new + LO 重算公式", xlsx_recalc)

    def xls_legacy():
        import xlrd
        x = work / "b.xlsx"
        xls = lo.convert(x, "xls", work)
        book = xlrd.open_workbook(str(xls))
        vals = book.sheet_by_index(0).row_values(0)
        assert "项目" in vals, vals
        return f"旧版 .xls 读回表头 {vals[:2]}"
    check("xlsx->xls + xlrd 读回", xls_legacy)

    # 4) pptx 生成 -> 缩略图
    pptx = work / "c.pptx"

    def pptx_build():
        writers.new_pptx({"title": "演示标题 Test", "subtitle": "office-docs",
                          "slides": [{"title": "第一页", "bullets": ["要点 A", "要点 B"],
                                      "notes": "备注：讲慢点"},
                                     {"title": "数据页", "table": {"rows": [["季度", "营收"], ["Q3", "1234"]]}}]},
                         pptx)
        txt = readers.text_of(pptx)
        for want in ("演示标题", "第一页", "要点 A", "1234"):
            assert want in txt, f"pptx 里找不到 {want!r}"
        return f"{len(txt)} 字符（封面/要点/表格都在）"
    check("pptx-new + 抽文本", pptx_build)

    def pptx_thumb():
        imgs = pdf_ops.to_images(lo.convert(pptx, "pdf", work), work / "thumbs", dpi=90)
        assert imgs and imgs[0].stat().st_size > 1000, imgs
        return f"{len(imgs)} 张 PNG，首张 {imgs[0].stat().st_size} B"
    check("pptx -> PDF -> PNG 缩略图", pptx_thumb)

    # 5) PDF 合并 / 拆分
    def pdf_merge_split():
        p1 = lo.convert(docx, "pdf", work / "one")
        p2 = lo.convert(work / "b.xlsx", "pdf", work / "two")
        merged = pdf_ops.merge([p1, p2], work / "merged.pdf")
        n = pdf_ops.info(merged)["pages"]
        parts = pdf_ops.split(merged, work / "parts", pages_per_file=1)
        assert n >= 2 and len(parts) == n, (n, len(parts))
        return f"合并 {n} 页，拆成 {len(parts)} 个文件"
    check("pdf-merge + pdf-split", pdf_merge_split)

    # 6) OCR：把 PDF 渲染成图片再包回 PDF，模拟扫描件
    def ocr_scan():
        import pymupdf
        src_pdf = lo.convert(docx, "pdf", work)
        png = pdf_ops.to_images(src_pdf, work / "render", dpi=200)[0]
        scan = work / "scan.pdf"
        with pymupdf.open() as doc:
            pix = pymupdf.open(str(png))
            rect = pix[0].rect
            page = doc.new_page(width=rect.width, height=rect.height)
            page.insert_image(rect, filename=str(png))
            doc.save(str(scan))
        txt = pdf_ops.ocr(scan, lang="chi_sim+eng", dpi=200)
        assert "中文" in txt or CN in txt, f"OCR 没认出中文：{txt[:120]!r}"
        return f"OCR 取回 {len(txt)} 字符，含中文"
    check("pdf-ocr（图片型 PDF，chi_sim）", ocr_scan)

    # 7) PDF -> Word
    def pdf_to_docx():
        src_pdf = lo.convert(docx, "pdf", work)
        out = pdf_ops.to_docx(src_pdf, work / "back.docx")
        txt = readers.text_of(out)
        assert "第一章" in txt or CN in txt, txt[:200]
        return f"{out.name} 取回 {len(txt)} 字符（保真度有限，需人工核）"
    check("pdf-to-docx（pdf2docx）", pdf_to_docx)

    # 8) 结构信息
    check("info/kind 识别", lambda: json.dumps(
        {p.name: readers.kind_of(p) for p in (docx, pptx, work / "b.xlsx")}, ensure_ascii=False))

    # 9) PDF 旋转 / 水印
    def pdf_rotate():
        from pypdf import PdfReader
        out = extras.rotate(work / "merged.pdf", 90, work / "rot.pdf", pages=[1])
        r = PdfReader(str(out))
        assert r.pages[0].rotation == 90, r.pages[0].rotation
        assert len(r.pages) == 2 and r.pages[1].rotation == 0
        return "第 1 页 90°、第 2 页不动"
    check("pdf-rotate", pdf_rotate)

    def pdf_watermark():
        out = extras.watermark(lo.convert(docx, "pdf", work), "机密 CONFIDENTIAL",
                               work / "wm.pdf", angle=45, fontsize=40)
        txt = readers.text_of(out)
        assert "CONFIDENTIAL" in txt or "机密" in txt, f"水印没进 PDF：{txt[:120]!r}"
        return "水印文字已在 PDF 文本层里（中文字体走内置 china-s）"
    check("pdf-watermark（含中文）", pdf_watermark)

    # 10) PDF 加密 / 解密
    def pdf_crypt():
        enc = extras.encrypt(work / "merged.pdf", "pw-123", work / "enc.pdf")
        from pypdf import PdfReader
        r = PdfReader(str(enc))
        assert r.is_encrypted, "产物没加密"
        assert r.decrypt("pw-123"), "正确密码解不开"
        assert not r.decrypt("wrong-pw"), "错误密码竟然能解开"
        dec = extras.decrypt(enc, "pw-123", work / "dec.pdf")
        assert len(PdfReader(str(dec)).pages) == 2
        return "AES-256 加密；错误密码被拒；解密后页数一致"
    check("pdf-encrypt / pdf-decrypt", pdf_crypt)

    # 11) PDF 表单：造一个真表单，列表 + 填值 + 回读
    def pdf_forms():
        import pymupdf
        form = work / "form.pdf"
        with pymupdf.open() as doc:
            page = doc.new_page()
            w = pymupdf.Widget()
            w.field_name, w.field_type = "name", pymupdf.PDF_WIDGET_TYPE_TEXT
            w.rect = pymupdf.Rect(72, 72, 300, 104)
            page.add_widget(w)
            doc.save(str(form))
        fields = extras.form_fields(form)
        assert "name" in fields, fields
        res = extras.form_fill(form, {"name": "张三 Zhang"}, work / "filled.pdf")
        assert not res["unknown_fields"], res
        from pypdf import PdfReader
        got = (PdfReader(str(work / "filled.pdf")).get_fields() or {})["name"].value
        assert str(got) == "张三 Zhang", got
        return "字段已列出并成功填值回读"
    check("pdf-forms（列表 + 填值）", pdf_forms)

    # 12) Excel 图表：不能只看"保存成功"，要确认图表 XML 真在包里
    def xlsx_chart():
        out = extras.chart(work / "b.xlsx", "bar", "数据!B2:B3", "数据!A2:A3",
                           work / "chart.xlsx", title="自检图表")
        assert extras.has_chart(out), "zip 里没有 xl/charts/chart*.xml"
        from openpyxl import load_workbook
        assert len(load_workbook(str(out))["数据"]._charts) == 1
        return "柱状图已写进 xlsx（xl/charts/chart1.xml）"
    check("xlsx-chart", xlsx_chart)

    # 13) Word 修订/批注盘点：注入一条真修订再数
    def docx_review():
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        src = work / "review.docx"
        d = Document(str(docx))
        run = d.paragraphs[0].runs[0]
        ins = OxmlElement("w:ins")
        ins.set(qn("w:id"), "1")
        ins.set(qn("w:author"), "selftest")
        ins.set(qn("w:date"), "2026-09-30T00:00:00Z")
        run._element.addprevious(ins)
        ins.append(run._element)
        d.save(str(src))
        r = extras.review(src)
        assert r["insertions"] == 1, r
        return f"盘点：插入 {r['insertions']} / 删除 {r['deletions']} / 批注 {len(r['comments'])}"
    check("docx-review（修订盘点）", docx_review)

    # 14b) 表单模板逆向：规格渲染 ＋ 合成表单的"印刷/手写"切分（合成件不含任何真实隐私）
    def form_spec_render():
        import form_template as ft
        from docx import Document
        spec = {"page": {"margins_cm": {"top": 2.0, "bottom": 2.0, "left": 2.0, "right": 2.0}},
                "title": {"text": "某某登记表", "font": "黑体", "size": 22, "bold": True},
                "body": [{"table": {"widths_cm": [3.5, 13.0], "rows": [
                    {"height_cm": 2.0, "cells": [{"text": "姓名", "align": "center"}, {"text": ""}]},
                    {"height_cm": 2.0, "cells": [{"text": "住址", "align": "center"}, {"text": ""}]}]}},
                    {"p": {"text": "填表说明：用黑色钢笔填写。", "size": 10.5}}]}
        out = ft.render_spec(spec, work / "form-spec.docx")
        d = Document(str(out))
        txt = "\n".join(p.text for p in d.paragraphs) + \
              "\n".join(c.text for t in d.tables for r in t.rows for c in r.cells)
        assert "某某登记表" in txt and "填表说明" in txt, txt[:120]
        assert len(d.tables) == 1 and len(d.tables[0].rows) == 2
        return "规格 → 干净空白模板 docx（标题/表格/说明都在）"
    check("form-render（规格渲染模板）", form_spec_render)

    def form_verify_geometry():
        """回归：列宽必须钉死、整表必须落在一页、线位必须落在目标 cm 上。

        这条是踩坑记录转成的自检 —— 只写 cell.width 时 LibreOffice 会自动适配，
        分隔线跑到 11.05cm（要的是 5.50），还顺带溢出到第 2 页。
        """
        import form_template as ft
        spec = {"page": {"size": "A4", "margins_cm": {"top": 2.0, "bottom": 1.0, "left": 1.9, "right": 0.8}},
                "verify": {"pages": 1, "tol_cm": 0.12,
                           "h_cm": [3.60, 5.20, 13.80, 15.10], "v_cm": [1.90, 5.50, 20.20]},
                "title": {"text": "几何自检表", "font": "宋体", "size": 22, "bold": True, "align": "center",
                          "line_pt": 30, "space_after_pt": 15.35},
                "body": [{"table": {"widths_cm": [3.6, 14.7], "rows": [
                    {"height_cm": 1.6, "height_rule": "exact",
                     "cells": [{"text": "标签", "align": "center", "valign": "center", "font": "宋体", "size": 10.5},
                               {"text": "", "font": "宋体", "size": 10.5}]},
                    {"height_cm": 8.6, "height_rule": "atLeast",
                     "cells": [{"paras": [{"text": "住　所", "align": "center", "font": "宋体", "size": 12}],
                                "valign": "center"},
                               {"text": "正文内容" * 10, "font": "宋体", "size": 14}]},
                    {"height_cm": 1.3, "height_rule": "exact",
                     "cells": [{"text": "合并行", "font": "宋体", "size": 14, "merge_right": 1}]}]}}]}
        r = ft.verify_spec(spec, work / "form-verify.docx")
        assert r["pages"] == 1, f"溢出到 {r['pages']} 页"
        assert r["ok"], json.dumps({k: r[k] for k in ("h", "v", "h_max_delta", "v_max_delta")}, ensure_ascii=False)
        return f"1 页 · 横线最大差 {r['h_max_delta']}cm · 竖线最大差 {r['v_max_delta']}cm"
    check("form-verify（列宽钉死 + 线位自检）", form_verify_geometry)

    def font_aliases():
        """中文别名必须落到衬线宋体/cwTeX 仿宋，而不是一律兜到 Noto Sans。"""
        d = lo.doctor()
        bad = d.get("font_aliases_bad") or []
        assert not bad, f"这些字体名没落到预期字库：{bad} -> " + \
                        json.dumps({k: v['match'] for k, v in d['font_aliases'].items() if k in bad}, ensure_ascii=False)
        assert d["cjk_fonts"] >= 20, d["cjk_fonts"]
        return f"9 个中文字体名别名全部就位 · 中文字库 {d['cjk_fonts']} 套"
    check("中文字体别名（宋体→衬线，别兜到 Noto Sans）", font_aliases)

    def form_split_synthetic():
        import cv2
        import numpy as np
        import form_template as ft
        from PIL import Image, ImageDraw, ImageFont
        img = np.full((1400, 1000, 3), 246, np.uint8)                 # 纸
        for y in (200, 600, 1000):
            cv2.line(img, (60, y), (940, y), (70, 70, 70), 3)          # 印刷横线
        for x in (60, 400, 940):
            cv2.line(img, (x, 200), (x, 1000), (70, 70, 70), 3)        # 印刷竖线
        pil = Image.fromarray(cv2.cvtColor(img, cv2.COLOR_BGR2RGB))
        fnt = ImageFont.truetype("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", 34)
        ImageDraw.Draw(pil).text((80, 300), "姓名", font=fnt, fill=(25, 25, 25))
        img = cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)           # 印刷字（非手写）
        for cx, cy in ((200, 330), (250, 330), (820, 780)):            # 假"手写"：几段曲线
            pts = np.array([[cx + dx, cy + int(18 * np.sin(dx / 11))] for dx in range(0, 90, 6)], np.int32)
            cv2.polylines(img, [pts], False, (18, 18, 18), 5)
        prof = {"page": {"size": [1000, 1400]},
                "grid": {"h": [{"pos": y, "from": 60, "to": 940} for y in (200, 600, 1000)],
                         "v": [{"pos": x, "from": 200, "to": 1000} for x in (60, 400, 940)]},
                "fill_regions": [[120, 280, 390, 390], [700, 700, 945, 830]],
                "blocks": [], "protect_margin": 5}
        parts = ft.split_ink(img, prof)
        hw = parts["handwriting"]
        assert hw[290:370, 150:380].sum() > 0, "填写区内的假手写没被切出来"
        assert hw[280:360, 70:170].sum() < hw[290:370, 150:380].sum() / 4, "印刷字被误判成手写"
        g = cv2.cvtColor(ft.blank_raster(img, prof), cv2.COLOR_BGR2GRAY)
        assert (g[300:370, 150:380] > 200).mean() > 0.95, "空白里手写没被抹掉"
        assert (g[290:340, 80:170] < 150).mean() > 0.02, "印刷字被误抹"
        return "合成表单：假手写被切出并抹掉，印刷字与表格线保住"
    check("form-extract / form-blank（合成表单）", form_split_synthetic)

    # 14) 中文公文版式（GB/T 9704-2012）
    def gongwen_layout():
        from docx import Document
        from docx.oxml.ns import qn
        spec = {"title": "关于开展文档格式自查的通知", "recipient": "各科室：",
                "blocks": [{"h1": "总体要求"}, {"p": "现就有关事项通知如下。"},
                           {"h2": "主要任务"}, {"h3": "时间安排"}, {"h4": "报送方式"}],
                "signature": {"org": "XX市人民政府办公室", "date": "2026年9月30日"},
                "note": "（此件公开发布）"}
        out = cn_profiles.build_docx(spec, work / "gongwen.docx", "gongwen")
        d = Document(str(out))
        s = d.sections[0]
        got = [round(s.top_margin.cm, 2), round(s.bottom_margin.cm, 2),
               round(s.left_margin.cm, 2), round(s.right_margin.cm, 2)]
        assert got == [3.7, 3.5, 2.8, 2.6], got
        assert round(s.page_width.mm) == 210 and round(s.page_height.mm) == 297
        assert round(s.footer_distance.cm, 2) == 2.8
        ps = [p for p in d.paragraphs if p.runs]
        assert ps[0].runs[0].font.name == "方正小标宋简体" and ps[0].runs[0].font.size.pt == 22, ps[0].text
        body = ps[3]
        assert body.runs[0].font.name == "仿宋_GB2312" and body.runs[0].font.size.pt == 16, body.text
        assert abs(body.paragraph_format.line_spacing.pt - 29.0) < 0.01, body.paragraph_format.line_spacing
        ind = body._p.pPr.ind
        assert ind.get(qn("w:firstLineChars")) == "200", ind.attrib
        fonts = [(p.text[:3], p.runs[0].font.name) for p in ps
                 if p.text[:2] in ("一、", "（一", "1.", "（1")]
        assert [f for _, f in fonts] == ["黑体", "楷体_GB2312", "仿宋_GB2312", "仿宋_GB2312"], fonts
        assert "PAGE" in s.footer.paragraphs[0]._p.xml, "页脚缺 PAGE 域"
        return "边距 3.7/3.5/2.8/2.6 · 三号仿宋 · 29磅固定行距 · 黑/楷/仿层次 · 页码域"
    check("docx-new --profile gongwen（版式）", gongwen_layout)

    def gongwen_render():
        out = cn_profiles.build_docx({"title": "格式自查通知", "recipient": "各科室：",
                                      "blocks": [{"h1": "总体要求"}, {"p": "正文一句。"},
                                                 {"h2": "任务"}, {"p": "另一句。"}]},
                                     work / "gongwen.pdf.docx", "gongwen", portable=True)
        pdf = lo.convert(out, "pdf", work)
        txt = readers.text_of(pdf)
        for want in ("格式自查通知", "一、总体要求", "（一）任务", "— 1 —"):
            assert want in txt, f"PDF 里缺 {want!r}：{txt[:200]!r}"
        from docx import Document
        assert Document(str(out)).paragraphs[1].runs[0].font.name.startswith("Noto"), "portable 未换字体"
        return "公文转 PDF 成功：层次序号自动编号、页码一字线、portable 走 Noto"
    check("gongwen 渲染 + --portable", gongwen_render)

    # 15) 合同 / 报告版式
    def contract_report():
        from docx import Document
        ct = cn_profiles.build_docx(
            {"title": "技术服务合同", "blocks": [{"h1": "服务内容"}, {"p": "甲方委托乙方……"},
                                                {"h1": "费用与支付"}, {"p": "……"}]},
            work / "contract.docx", "contract")
        rp = cn_profiles.build_docx(
            {"title": "2026年第三季度经营分析报告", "blocks": [{"h1": "总体情况"}, {"p": "……"}]},
            work / "report.docx", "report")
        d = Document(str(ct))
        ps = [p for p in d.paragraphs if p.runs]
        assert round(d.sections[0].left_margin.cm, 2) == 3.17
        assert ps[0].runs[0].font.name == "黑体" and ps[0].runs[0].font.size.pt == 22
        assert ps[4].runs[0].font.name == "宋体" and ps[4].runs[0].font.size.pt == 12
        assert abs(ps[4].paragraph_format.line_spacing - 1.5) < 0.01
        assert any(p.text.startswith("第一条") for p in ps) and any(p.text.startswith("第二条") for p in ps)
        assert Document(str(rp)).paragraphs[0].runs[0].font.size.pt == 18
        return "合同：第一条/第二条自动编号、小四宋体 1.5 倍；报告：小二标题"
    check("docx-new --profile contract / report", contract_report)

    # 16) 汇报 PPT（brief）
    def brief_pptx():
        from pptx import Presentation
        out = cn_profiles.build_pptx(
            {"title": "Q3 经营汇报", "subtitle": "2026-09-30",
             "slides": [{"title": "结论", "bullets": [f"要点{i}" for i in range(1, 9)]},
                        {"title": "数据", "table": {"rows": [["季度", "营收"], ["Q3", "1234"]]}}]},
            work / "brief.pptx", "brief")
        prs = Presentation(str(out))
        assert len(prs.slides) == 3
        assert round(prs.slide_width / 914400, 2) == 13.33, prs.slide_width
        cover = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
        assert cover.font.size.pt == 32 and cover.font.name == "微软雅黑", (cover.font.size.pt, cover.font.name)
        body = prs.slides[1].shapes[1].text_frame.paragraphs[0].runs[0]
        assert body.font.size.pt == 18
        bullets = prs.slides[1].shapes[1].text_frame.paragraphs
        assert len(bullets) == 7 and "已截断" in bullets[-1].runs[0].text, len(bullets)
        return "16:9 · 封面32pt/正文18pt · 单页要点自动截到 6 条"
    check("pptx-new --profile brief", brief_pptx)

    # 17) docx 跨 run 查找替换（Word 把一句话拆成多个 run 时会漏词）
    def docx_cross_run_replace():
        from docx import Document
        from docx.oxml.ns import qn
        src = work / "frag.docx"
        d = Document()
        p = d.add_paragraph()
        for t in ("北京", "市朝", "阳区"):
            el = p.add_run("")._element
            te = el.makeelement(qn("w:t"), {})
            te.text = t
            el.append(te)
        d.save(str(src))
        d2 = Document(str(src))
        assert len(d2.paragraphs[0].runs) == 3, "构造失败：run 没被拆开"
        assert all("北京市" not in r.text for r in d2.paragraphs[0].runs), "构造失败：词没跨 run"
        res = writers.docx_replace(src, {"北京市": "上海市"}, work / "frag-out.docx")
        assert res["hits"] == 1, res
        got = Document(str(work / "frag-out.docx")).paragraphs[0].text
        assert got == "上海市朝阳区", got
        return f"词跨 3 个 run，逐 run 找必漏；渲染文本匹配 → {got}"

    check("docx-replace 跨 run 命中", docx_cross_run_replace)

    def docx_merge_runs_case():
        import docx_runs as dr
        res = dr.merge_runs_in_file(work / "frag.docx", work / "frag-merged.docx")
        assert res["runs_after"] < res["runs_before"], res
        from docx import Document
        assert Document(str(work / "frag-merged.docx")).paragraphs[0].text == "北京市朝阳区"
        return f"{res['runs_before']} 个 run → {res['runs_after']} 个，文字不变"

    check("docx-merge-runs 合并碎片", docx_merge_runs_case)

    # 18) xlsx 公式重算（openpyxl 写公式不算值，交付前必须重算）
    def xlsx_recalc_case():
        import xlsx_recalc as xr
        from openpyxl import Workbook, load_workbook
        f = work / "calc.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "数据"
        ws["A1"], ws["A2"], ws["A3"] = 1, 2, 3
        ws["B1"], ws["B2"], ws["C1"] = "=SUM(A1:A3)", "=B1*10", "=1/0"
        wb.save(str(f))
        assert load_workbook(f, data_only=True)["数据"]["B1"].value is None, \
            "openpyxl 竟然写了缓存值（前提变了，重算的必要性要重估）"
        r = xr.recalc(f, timeout=120)
        assert r["status"] == "errors_found", r
        w = load_workbook(f, data_only=True)["数据"]
        assert (w["B1"].value, w["B2"].value) == (6, 60), (w["B1"].value, w["B2"].value)
        assert "#DIV/0!" in str(w["C1"].value), w["C1"].value
        return "重算前 B1=None → 重算后 B1=6、B2=60，并抓到 C1 的 #DIV/0!"

    check("xlsx-recalc 重算公式 + 扫错误", xlsx_recalc_case)

    # 19) OOXML 包体检（Word「无法读取的内容」的硬伤）
    def ooxml_gate_case():
        import ooxml_check
        import zipfile as _zip
        good = ooxml_check.check(docx)
        assert good["ok"] and good["kind"] == "docx", good
        broken = work / "broken.docx"
        with _zip.ZipFile(docx) as zin, _zip.ZipFile(broken, "w", _zip.ZIP_DEFLATED) as zout:
            for n in zin.namelist():
                if n == "word/styles.xml":
                    continue                      # 部件删了，Content_Types 的 Override 还留着
                zout.writestr(n, zin.read(n))
        bad = ooxml_check.check(broken)
        assert not bad["ok"] and any("styles.xml" in e for e in bad["errors"]), bad
        return f"好包 {good['stats']['parts']} 部件通过；删 styles.xml 后报 {len(bad['errors'])} 条"

    check("ooxml-check 包结构闸门", ooxml_gate_case)

    # 20) 静态 PDF：抽结构 → 框位自检 → 按框填字 → 出校验图
    def pdf_form_fill_case():
        import pymupdf
        import pdf_forms as pf
        f = work / "static-form.pdf"
        doc = pymupdf.open()
        page = doc.new_page(width=595, height=842)
        for y in (100, 140):
            page.draw_line(pymupdf.Point(60, y), pymupdf.Point(535, y), width=0.8)
        page.insert_text((70, 126), "姓名", fontname="china-s", fontsize=11)
        doc.save(str(f))
        doc.close()

        st = pf.extract_structure(f)["pages"][0]
        assert any(l["text"] == "姓名" for l in st["labels"]), st["labels"]
        assert len(st["h_lines"]) == 2, st["h_lines"]

        spec = {"unit": "pt", "pages": [{"page": 1, "size_pt": [595, 842]}],
                "fields": [{"page": 1, "label": "姓名", "box": [170, 104, 300, 136],
                            "text": "张三丰", "font_size": 11}]}
        assert pf.check_boxes(spec)["ok"], pf.check_boxes(spec)
        out = work / "static-form-filled.pdf"
        png = work / "static-form-preview.png"
        res = pf.fill(f, spec, out, preview=png)
        assert res["ok"] and res["overflow"] == [], res
        assert "张三丰" in readers.text_of(out)
        assert png.exists() and png.stat().st_size > 1000
        # 自检必须挡得住装不下的框（先拦，而不是填完才发现）
        bad = {"unit": "pt", "fields": [{"label": "小框", "box": [10, 10, 40, 22],
                                         "text": "这一段字肯定装不下", "font_size": 12}]}
        assert not pf.check_boxes(bad)["ok"], "小框没被拦住"
        return "抽到 1 标签/2 横线 → 自检过 → 填字进 PDF → 出校验图；装不下的小框被拦"

    check("pdf-structure / pdf-fill / 框位自检", pdf_form_fill_case)

    ok = sum(1 for _, good, _ in RESULTS if good)
    bad = [(n, d) for n, good, d in RESULTS if not good]
    print("\n" + "=" * 60)
    print(f"自检：{ok}/{len(RESULTS)} 通过")
    for n, d in bad:
        print(f"  FAIL {n}: {d}")
    print(f"产物目录（已保留供人工看）：{work}")
    return 0 if not bad else 1


if __name__ == "__main__":
    sys.exit(main())
