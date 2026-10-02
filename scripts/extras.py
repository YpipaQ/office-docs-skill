#!/usr/bin/env python3
"""补齐常用但少见的文档操作：PDF 旋转/水印/加密/表单，Excel 图表，Word 修订与批注盘点。

全部用公开 API 自己实现（pypdf / PyMuPDF / openpyxl / python-docx），不搬运任何第三方技能材料。
"""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

CJK_PDF_FONT = "china-s"   # PyMuPDF 内置简体中文字体；ASCII 场景回退 helv


# ---------------- PDF ----------------

def rotate(src: str | Path, angle: int, out: str | Path | None = None,
           pages: list[int] | None = None) -> Path:
    """顺时针旋转整份或指定页（1-based）；角度必须是 90 的倍数。"""
    from pypdf import PdfReader, PdfWriter
    if angle % 90 != 0:
        raise ValueError("角度必须是 90 的倍数")
    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-rot{angle}.pdf")
    reader, writer = PdfReader(str(src)), PdfWriter()
    targets = {p - 1 for p in pages} if pages else set(range(len(reader.pages)))
    for i, page in enumerate(reader.pages):
        if i in targets:
            page.rotate(angle)
        writer.add_page(page)
    with out.open("wb") as fh:
        writer.write(fh)
    return out


def watermark(src: str | Path, text: str, out: str | Path | None = None, *,
              angle: int = 45, fontsize: int = 60, opacity: float = 0.12,
              color: tuple[float, float, float] = (0.5, 0.5, 0.5)) -> Path:
    """给每页盖一层浅色文字水印（PyMuPDF 造戳，pypdf 合并），CJK 走内置 china-s 字体。"""
    import pymupdf
    from pypdf import PdfReader, PdfWriter

    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-wm.pdf")
    fontname = CJK_PDF_FONT if any(ord(c) > 127 for c in text) else "helv"

    with pymupdf.open(str(src)) as doc:
        w, h = doc[0].rect.width, doc[0].rect.height
    stamp_doc = pymupdf.open()
    page = stamp_doc.new_page(width=w, height=h)
    pivot = pymupdf.Point(w / 2, h / 2)
    try:
        page.insert_text(pivot, text, fontsize=fontsize, fontname=fontname,
                         color=color, fill_opacity=opacity,
                         morph=(pivot, pymupdf.Matrix(angle)))
    except Exception:
        # 个别字体/旋转组合不被支持：退化为水平居中，宁可有水印也不要失败
        page.insert_text((w * 0.12, h / 2), text, fontsize=fontsize,
                         fontname=fontname, color=color, fill_opacity=opacity)
    stamp_bytes = stamp_doc.tobytes()
    stamp_doc.close()

    stamp_page = PdfReader(__import__("io").BytesIO(stamp_bytes)).pages[0]
    reader, writer = PdfReader(str(src)), PdfWriter()
    for p in reader.pages:
        p.merge_page(stamp_page)
        writer.add_page(p)
    with out.open("wb") as fh:
        writer.write(fh)
    return out


def encrypt(src: str | Path, password: str, out: str | Path | None = None, *,
            owner_password: str | None = None, algorithm: str = "AES-256") -> Path:
    from pypdf import PdfReader, PdfWriter
    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-enc.pdf")
    reader, writer = PdfReader(str(src)), PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, owner_password=owner_password or password,
                   algorithm=algorithm)
    with out.open("wb") as fh:
        writer.write(fh)
    return out


def decrypt(src: str | Path, password: str, out: str | Path | None = None) -> Path:
    from pypdf import PdfReader, PdfWriter
    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-dec.pdf")
    reader = PdfReader(str(src))
    if reader.is_encrypted and not reader.decrypt(password):
        raise ValueError("密码不对，解不开")
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    with out.open("wb") as fh:
        writer.write(fh)
    return out


def form_fields(src: str | Path) -> dict:
    """列出 AcroForm 字段（名字、类型、当前值、可选项）。"""
    from pypdf import PdfReader
    reader = PdfReader(str(Path(src).expanduser().resolve()))
    if reader.is_encrypted:
        reader.decrypt("")
    fields = reader.get_fields() or {}
    out: dict = {}
    for name, f in fields.items():
        states = getattr(f, "states", None) or getattr(f, "_States_", None)
        out[name] = {
            "type": str(getattr(f, "field_type", getattr(f, "/FT", ""))),
            "value": str(getattr(f, "value", getattr(f, "/V", "")) or ""),
            "options": sorted(states.keys()) if isinstance(states, dict) else None,
        }
    return out


def _ensure_default_resources(writer) -> None:
    """给 AcroForm 补一个最小 `/DR`（Helvetica）＋ `/DA`。

    pypdf 的 `update_page_form_field_values` 在 `/DR` 缺失时会把条目当普通 dict 用，直接崩
    （`'dict' object has no attribute 'get_object'`）。PyMuPDF 生成的表单正是这种，Acrobat 表单一般自带。
    """
    from pypdf.generic import DictionaryObject, NameObject, TextStringObject
    root = writer._root_object
    if "/AcroForm" not in root:
        root[NameObject("/AcroForm")] = DictionaryObject()
    acro = root["/AcroForm"].get_object()
    if "/DR" not in acro:
        acro[NameObject("/DR")] = DictionaryObject()
    dr = acro["/DR"].get_object()
    if "/Font" not in dr:
        helv = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        dr[NameObject("/Font")] = DictionaryObject({NameObject("/Helv"): writer._add_object(helv)})
    if "/DA" not in acro:
        acro[NameObject("/DA")] = TextStringObject("/Helv 0 Tf 0 g")


def form_fill(src: str | Path, values: dict[str, str], out: str | Path | None = None) -> dict:
    """填表单。对象是 AcroForm 文本框/勾选框；值必须与字段名一一对上。"""
    from pypdf import PdfReader, PdfWriter
    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-filled.pdf")
    reader, writer = PdfReader(str(src)), PdfWriter()
    if reader.is_encrypted:
        reader.decrypt("")
    known = set((reader.get_fields() or {}).keys())
    unknown = [k for k in values if k not in known]
    writer.append(reader)
    _ensure_default_resources(writer)
    if hasattr(writer, "set_need_appearances_writer"):
        writer.set_need_appearances_writer(True)
    for page in writer.pages:
        writer.update_page_form_field_values(page, values, auto_regenerate=False)
    with out.open("wb") as fh:
        writer.write(fh)
    return {"out": str(out), "filled": sorted(set(values) - set(unknown)), "unknown_fields": unknown}


# ---------------- XLSX 图表 ----------------

def chart(src: str | Path, chart_type: str, data_ref: str, cats_ref: str | None = None,
          out: str | Path | None = None, *, title: str | None = None,
          anchor: str = "E2", series_titles: bool = True) -> Path:
    """给现成 xlsx 加图表。引用写法 `Sheet1!B2:B5`；title_from_data 表头当系列名。"""
    from openpyxl import load_workbook
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.utils.cell import range_boundaries

    def _ref(ref: str):
        if "!" in ref:
            sheet, rng = ref.split("!", 1)
        else:
            sheet, rng = wb.sheetnames[0], ref
        min_col, min_row, max_col, max_row = range_boundaries(rng)
        return sheet, Reference(ws_of(sheet), min_col=min_col, min_row=min_row,
                                max_col=max_col, max_row=max_row)

    src = Path(src).expanduser().resolve()
    out = Path(out).expanduser() if out else src.with_name(f"{src.stem}-chart.xlsx")
    wb = load_workbook(str(src))
    ws_of = lambda name: wb[name] if name in wb.sheetnames else wb.worksheets[0]
    cls = {"bar": BarChart, "line": LineChart, "pie": PieChart}.get(chart_type.lower())
    if cls is None:
        raise ValueError("图表类型只支持 bar / line / pie")
    ch = cls()
    ch.title = title or f"{chart_type.upper()} 图"
    sheet, data = _ref(data_ref)
    cats = None
    if cats_ref:
        _, cats = _ref(cats_ref)
    if cats is not None:
        ch.set_categories(cats)
    ch.add_data(data, titles_from_data=series_titles)
    wb[sheet].add_chart(ch, anchor)
    wb.save(str(out))
    return out


def has_chart(src: str | Path) -> bool:
    """产物里是否真有图表 XML（自检用，避免只信"保存成功"）。"""
    with zipfile.ZipFile(Path(src).expanduser()) as z:
        return any(n.startswith("xl/charts/chart") for n in z.namelist())


# ---------------- DOCX 修订 / 批注 ----------------

def review(src: str | Path) -> dict:
    """盘点 Word 的修订与批注（只读）：插入/删除/格式变更条数 + 批注作者与正文。"""
    from docx import Document
    src = Path(src).expanduser().resolve()
    xml = Document(str(src)).element.body.xml
    out = {
        "insertions": len(re.findall(r"<w:ins[ >]", xml)),
        "deletions": len(re.findall(r"<w:del[ >]", xml)),
        "format_changes": len(re.findall(r"<w:rPrChange[ >]", xml)),
        "comments": [],
        "comment_replies": 0,
    }
    with zipfile.ZipFile(src) as z:
        name = "word/comments.xml"
        if name in z.namelist():
            cx = z.read(name).decode("utf-8", "replace")
            for author, _id, text in re.findall(
                    r'<w:comment [^>]*w:author="([^"]*)"[^>]*w:id="([^"]*)"[^>]*>(.*?)</w:comment>', cx, re.S):
                out["comments"].append({"author": author, "id": _id,
                                        "text": re.sub(r"<[^>]+>", "", text).strip()})
            out["comment_replies"] = len(re.findall(r'<w:comment [^>]*w:parent="', cx))
    return out
