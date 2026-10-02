#!/usr/bin/env python3
"""统一读取：抽文本 + 看结构。优先用原生库（快、不改文件），兜底走 LibreOffice 转换。

支持：docx / doc / xlsx / xls / ods / pptx / ppt / odp / pdf / csv / tsv / txt / md。
"""
from __future__ import annotations

import csv
import io
import zipfile
from pathlib import Path

import convert as lo

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm",
            ".log", ".py", ".sh", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".rst", ".tex"}


def _looks_textual(p: Path) -> bool:
    """未知扩展名时，用"能不能 UTF-8 解码且没有 NUL"判断是不是纯文本，避免把文本塞给 LibreOffice。"""
    try:
        chunk = p.open("rb").read(65536)
    except OSError:
        return False
    if b"\x00" in chunk:
        return False
    try:
        chunk.decode("utf-8")
        return True
    except UnicodeDecodeError:
        return False


def kind_of(path: str | Path) -> str:
    """按内容而不是扩展名判断真实类型（扩展名经常是骗人的）。"""
    p = Path(path)
    ext = p.suffix.lower()
    try:
        head = p.open("rb").read(8)
    except OSError:
        return "missing"
    if head.startswith(b"%PDF"):
        return "pdf"
    if head[:4] == b"PK\x03\x04":
        try:
            with zipfile.ZipFile(p) as z:
                names = set(z.namelist())
            for marker, k in (("word/document.xml", "docx"), ("xl/workbook.xml", "xlsx"),
                              ("ppt/presentation.xml", "pptx"), ("content.xml", "odf")):
                if marker in names:
                    return k
        except zipfile.BadZipFile:
            return "zip"
        return "zip"
    if head[:4] == b"\xd0\xcf\x11\xe0":  # OLE2 复合文档
        return {".doc": "doc", ".xls": "xls", ".ppt": "ppt"}.get(ext, "ole2")
    if ext in TEXT_EXT:
        return ext.lstrip(".")
    return ext.lstrip(".") or "unknown"


def text_of(path: str | Path, *, max_chars: int | None = None) -> str:
    p = Path(path).expanduser().resolve()
    if not p.exists():
        raise FileNotFoundError(p)
    k = kind_of(p)
    parts: list[str] = []

    if k == "docx":
        from docx import Document
        d = Document(str(p))
        for para in d.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for t in d.tables:
            for row in t.rows:
                parts.append("\t".join(c.text.strip() for c in row.cells))
    elif k == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=False, read_only=True)
        for ws in wb.worksheets:
            parts.append(f"# sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() for c in row):
                    parts.append("\t".join("" if c is None else str(c) for c in row))
        wb.close()
    elif k == "pptx":
        from pptx import Presentation
        prs = Presentation(str(p))
        for i, slide in enumerate(prs.slides, 1):
            parts.append(f"# slide {i}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
                if getattr(shape, "has_table", False) and shape.has_table:
                    for row in shape.table.rows:
                        parts.append("\t".join(c.text.strip() for c in row.cells))
    elif k == "pdf":
        try:
            import pymupdf
            with pymupdf.open(str(p)) as doc:
                parts = [page.get_text() for page in doc]
        except Exception:
            from pypdf import PdfReader
            parts = [(pg.extract_text() or "") for pg in PdfReader(str(p)).pages]
    elif k == "xls":
        import xlrd
        book = xlrd.open_workbook(str(p))
        for sheet in book.sheets():
            parts.append(f"# sheet: {sheet.name}")
            for r in range(sheet.nrows):
                parts.append("\t".join(str(v) for v in sheet.row_values(r)))
    elif k == "odf":
        from odf import teletype
        from odf.opendocument import load
        from odf.text import P
        doc = load(str(p))
        parts = [teletype.extractText(p) for p in doc.getElementsByType(P)]
    elif k in {"csv", "tsv", "txt", "md"}:
        return p.read_text(encoding="utf-8", errors="replace")[:max_chars] if max_chars else \
            p.read_text(encoding="utf-8", errors="replace")
    else:
        if _looks_textual(p):  # 未知扩展名但确实是文本 -> 直接读，别惊动 LibreOffice
            txt = p.read_text(encoding="utf-8", errors="replace")
            return txt[:max_chars] if max_chars else txt
        # 真兜底：让 LibreOffice 自己导出成纯文本
        out = lo.convert(p, "txt")
        return out.read_text(encoding="utf-8", errors="replace")

    out = "\n".join(x for x in parts if x and x.strip())
    return out[:max_chars] if max_chars else out


def info_of(path: str | Path) -> dict:
    p = Path(path).expanduser().resolve()
    k = kind_of(p)
    d: dict = {"path": str(p), "kind": k, "bytes": p.stat().st_size if p.exists() else 0}
    try:
        if k == "docx":
            from docx import Document
            doc = Document(str(p))
            d.update(paragraphs=len(doc.paragraphs), tables=len(doc.tables),
                     sections=len(doc.sections),
                     words=sum(len(x.text.split()) for x in doc.paragraphs))
        elif k == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(p), read_only=True)
            d.update(sheets=[ws.title for ws in wb.worksheets],
                     dimensions={ws.title: ws.calculate_dimension() for ws in wb.worksheets})
            wb.close()
        elif k == "pptx":
            from pptx import Presentation
            prs = Presentation(str(p))
            d.update(slides=len(prs.slides), slide_size=[prs.slide_width, prs.slide_height])
        elif k == "pdf":
            import pymupdf
            with pymupdf.open(str(p)) as doc:
                d.update(pages=doc.page_count, encrypted=doc.is_encrypted,
                         meta={k2: v for k2, v in (doc.metadata or {}).items() if v})
        elif k in {"doc", "xls", "ppt", "ole2"}:
            d.update(note="旧版二进制格式；文本/转换请走 LibreOffice")
    except Exception as exc:  # 结构信息拿不到不该让整个调用失败
        d["error"] = f"{type(exc).__name__}: {exc}"
    return d


def read_table(path: str | Path) -> list[list[str]]:
    """把一个工作表/CSV 读成二维字符串表（CSV/TSV 用标准库，表格文件用 openpyxl）。"""
    p = Path(path).expanduser().resolve()
    k = kind_of(p)
    if k in {"csv", "tsv"}:
        delim = "\t" if k == "tsv" else ","
        with p.open(encoding="utf-8-sig", errors="replace", newline="") as fh:
            return [list(r) for r in csv.reader(fh, delimiter=delim)]
    if k == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(p), data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c) for c in row] for row in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    raise ValueError(f"{p.name}: 不支持直接读表（kind={k}）")


def write_csv(rows: list[list[str]], out: str | Path) -> Path:
    out = Path(out).expanduser()
    out.parent.mkdir(parents=True, exist_ok=True)
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    out.write_text(buf.getvalue(), encoding="utf-8")  # 中文一律 UTF-8，别用系统默认
    return out
